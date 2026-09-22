"""
PDF Analyzer — одностраничное приложение Streamlit.

Пайплайн: загрузка PDF -> извлечение текста (pypdf) -> асинхронный анализ
через облачный LLM API (OpenAI-совместимый) -> st.dataframe -> экспорт в Excel (openpyxl).

Запуск:       streamlit run app.py
Зависимости:  pip install streamlit pypdf httpx pandas openpyxl
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
from datetime import datetime
from typing import Any, Callable

import httpx
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from pypdf.errors import PdfReadError

# ----------------------------------------------------------------------------
# Конфигурация
# ----------------------------------------------------------------------------
REQUEST_TIMEOUT = 90.0  # секунд на один запрос к API

PROVIDERS: dict[str, dict[str, Any]] = {
    "OpenAI": {
        "base_url": "https://api.openai.com/v1",
        "models": ["gpt-4o-mini", "gpt-4o", "gpt-4.1-mini"],
    },
    "OpenRouter": {
        "base_url": "https://openrouter.ai/api/v1",
        "models": [
            "openai/gpt-4o-mini",
            "anthropic/claude-3.5-haiku",
            "google/gemini-flash-1.5",
            "meta-llama/llama-3.1-70b-instruct",
        ],
    },
    "Groq": {
        "base_url": "https://api.groq.com/openai/v1",
        "models": ["llama-3.3-70b-versatile", "llama-3.1-8b-instant", "mixtral-8x7b-32768"],
    },
    "Свой endpoint": {"base_url": "", "models": []},
}

SYSTEM_PROMPT = (
    "Ты — аналитик документов. Проанализируй текст и верни СТРОГО валидный JSON "
    "без пояснений и markdown, со следующими полями:\n"
    '{"topic": "тема документа (до 10 слов)", '
    '"summary": "краткое содержание (2-3 предложения)", '
    '"keywords": ["до 7 ключевых слов"], '
    '"doc_type": "тип документа (договор, статья, отчёт, инструкция и т.п.)", '
    '"language": "язык документа"}'
)

RESULT_COLUMNS = [
    "Файл", "Страниц", "Символов", "Тема", "Краткое содержание",
    "Ключевые слова", "Тип документа", "Язык", "Модель", "Статус", "Ошибка", "Время",
]


# ----------------------------------------------------------------------------
# Состояние
# ----------------------------------------------------------------------------
def init_state() -> None:
    """Инициализирует ключи st.session_state один раз за сессию."""
    st.session_state.setdefault("results", [])        # list[dict] — строки таблицы
    st.session_state.setdefault("processed", set())   # fingerprint'ы обработанных файлов
    st.session_state.setdefault("api_key", default_api_key())


def default_api_key() -> str:
    """Берёт ключ из st.secrets или переменной окружения, если они есть."""
    try:
        if "API_KEY" in st.secrets:
            return str(st.secrets["API_KEY"])
    except Exception:  # secrets.toml отсутствует — это нормально
        pass
    return os.getenv("API_KEY", "")


# ----------------------------------------------------------------------------
# Работа с PDF
# ----------------------------------------------------------------------------
def file_fingerprint(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def extract_text(data: bytes) -> tuple[str, int]:
    """
    Извлекает текст из PDF. Возвращает (текст, число_страниц).
    Бросает ValueError с человекочитаемым сообщением при проблемах.
    """
    try:
        reader = PdfReader(io.BytesIO(data))
    except PdfReadError as exc:
        raise ValueError(f"Файл повреждён или не является PDF: {exc}") from exc

    if reader.is_encrypted:
        try:
            reader.decrypt("")  # пробуем пустой пароль
        except Exception as exc:
            raise ValueError("PDF защищён паролем") from exc

    pages_text: list[str] = []
    for page in reader.pages:
        try:
            pages_text.append(page.extract_text() or "")
        except Exception:
            pages_text.append("")  # одна битая страница не должна ломать весь файл

    text = "\n".join(pages_text).strip()
    if not text:
        raise ValueError("Текст не найден (возможно, это скан — нужен OCR)")
    return text, len(reader.pages)


# ----------------------------------------------------------------------------
# Асинхронный анализ через API
# ----------------------------------------------------------------------------
def parse_model_json(content: str) -> dict[str, Any]:
    """Достаёт JSON из ответа модели, даже если он обёрнут в ```json ... ```."""
    start, end = content.find("{"), content.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("В ответе модели нет JSON-объекта")
    return json.loads(content[start : end + 1])


def make_row(item: dict[str, Any], model: str, **fields: Any) -> dict[str, Any]:
    """Формирует строку результата с единым набором колонок."""
    row = {col: "" for col in RESULT_COLUMNS}
    row.update(
        {
            "Файл": item["name"],
            "Страниц": item.get("pages", 0),
            "Символов": item.get("chars", 0),
            "Модель": model,
            "Статус": "OK",
            "Время": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
    )
    row.update(fields)
    return row


async def analyze_one(
    client: httpx.AsyncClient,
    semaphore: asyncio.Semaphore,
    item: dict[str, Any],
    model: str,
    max_chars: int,
) -> dict[str, Any]:
    """Отправляет текст одного файла в API и возвращает строку результата."""
    payload = {
        "model": model,
        "temperature": 0.2,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": item["text"][:max_chars]},
        ],
    }

    async with semaphore:
        try:
            resp = await client.post("/chat/completions", json=payload)
            resp.raise_for_status()
            content = resp.json()["choices"][0]["message"]["content"]
            data = parse_model_json(content)

            keywords = data.get("keywords", [])
            if isinstance(keywords, list):
                keywords = ", ".join(map(str, keywords))

            return make_row(
                item,
                model,
                **{
                    "Тема": data.get("topic", ""),
                    "Краткое содержание": data.get("summary", ""),
                    "Ключевые слова": keywords,
                    "Тип документа": data.get("doc_type", ""),
                    "Язык": data.get("language", ""),
                },
            )
        except httpx.HTTPStatusError as exc:
            error = f"HTTP {exc.response.status_code}: {exc.response.text[:200]}"
        except httpx.TimeoutException:
            error = f"Таймаут запроса ({REQUEST_TIMEOUT:.0f} с)"
        except httpx.RequestError as exc:
            error = f"Сетевая ошибка: {exc}"
        except (KeyError, IndexError, json.JSONDecodeError, ValueError) as exc:
            error = f"Некорректный ответ модели: {exc}"

    return make_row(item, model, **{"Статус": "Ошибка", "Ошибка": error})


async def analyze_all(
    items: list[dict[str, Any]],
    *,
    base_url: str,
    api_key: str,
    model: str,
    max_chars: int,
    concurrency: int,
    on_progress: Callable[[int, int], None],
) -> list[dict[str, Any]]:
    """Параллельно анализирует все файлы, ограничивая число одновременных запросов."""
    semaphore = asyncio.Semaphore(concurrency)
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}

    async with httpx.AsyncClient(
        base_url=base_url.rstrip("/"), headers=headers, timeout=REQUEST_TIMEOUT
    ) as client:
        tasks = [analyze_one(client, semaphore, item, model, max_chars) for item in items]
        results: list[dict[str, Any]] = []
        for coro in asyncio.as_completed(tasks):
            results.append(await coro)
            on_progress(len(results), len(tasks))
    return results


# ----------------------------------------------------------------------------
# Экспорт в Excel
# ----------------------------------------------------------------------------
def build_excel(df: pd.DataFrame) -> bytes:
    """Собирает .xlsx через openpyxl: заголовок, автоширина, фильтр, закреплённая шапка."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Анализ PDF"

    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in df.itertuples(index=False):
        ws.append(["" if pd.isna(v) else v for v in row])

    for idx, col in enumerate(df.columns, start=1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].tolist()])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(12, longest + 2), 60)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
def render_sidebar() -> dict[str, Any]:
    """Настройки API и модели. Возвращает словарь параметров."""
    st.sidebar.header("⚙️ Настройки")

    provider = st.sidebar.selectbox("Провайдер", list(PROVIDERS))
    cfg = PROVIDERS[provider]

    base_url = st.sidebar.text_input(
        "Base URL", value=cfg["base_url"], disabled=bool(cfg["base_url"])
    )

    if cfg["models"]:
        model = st.sidebar.selectbox("Модель", cfg["models"])
    else:
        model = st.sidebar.text_input("Модель", placeholder="например, llama3")

    st.session_state.api_key = st.sidebar.text_input(
        "API ключ", value=st.session_state.api_key, type="password"
    )

    max_chars = st.sidebar.slider(
        "Макс. символов в запросе", 2_000, 60_000, 12_000, step=1_000,
        help="Обрезает текст документа, чтобы уложиться в контекст модели",
    )
    concurrency = st.sidebar.slider("Параллельных запросов", 1, 10, 4)

    return {
        "base_url": base_url,
        "model": model,
        "api_key": st.session_state.api_key,
        "max_chars": max_chars,
        "concurrency": concurrency,
    }


def validate_settings(settings: dict[str, Any]) -> list[str]:
    errors = []
    if not settings["base_url"].startswith("http"):
        errors.append("Укажите корректный Base URL")
    if not settings["model"]:
        errors.append("Укажите модель")
    if not settings["api_key"]:
        errors.append("Укажите API ключ")
    return errors


def process_files(uploaded_files: list, settings: dict[str, Any]) -> None:
    """Извлекает текст, запускает асинхронный анализ и сохраняет результаты в состояние."""
    items: list[dict[str, Any]] = []
    new_rows: list[dict[str, Any]] = []
    skipped = 0

    # 1. Синхронный этап: извлечение текста (ошибки сразу попадают в таблицу)
    with st.spinner("Извлечение текста из PDF…"):
        for f in uploaded_files:
            data = f.getvalue()
            fp = file_fingerprint(data)
            if fp in st.session_state.processed:
                skipped += 1
                continue
            st.session_state.processed.add(fp)

            item = {"name": f.name, "fingerprint": fp}
            try:
                text, pages = extract_text(data)
                item.update(text=text, pages=pages, chars=len(text))
                items.append(item)
            except ValueError as exc:
                new_rows.append(
                    make_row(item, settings["model"], **{"Статус": "Ошибка", "Ошибка": str(exc)})
                )

    if skipped:
        st.info(f"Пропущено уже обработанных файлов: {skipped}")

    # 2. Асинхронный этап: запросы к API
    if items:
        progress = st.progress(0.0, text="Отправка в модель…")

        def on_progress(done: int, total: int) -> None:
            progress.progress(done / total, text=f"Обработано {done} из {total}")

        try:
            api_rows = asyncio.run(
                analyze_all(
                    items,
                    base_url=settings["base_url"],
                    api_key=settings["api_key"],
                    model=settings["model"],
                    max_chars=settings["max_chars"],
                    concurrency=settings["concurrency"],
                    on_progress=on_progress,
                )
            )
            new_rows.extend(api_rows)
        except Exception as exc:  # непредвиденный сбой всего пула
            st.error(f"Критическая ошибка анализа: {exc}")
            # откатываем отметку «обработано», чтобы можно было повторить
            for item in items:
                st.session_state.processed.discard(item["fingerprint"])
        finally:
            progress.empty()

    st.session_state.results.extend(new_rows)

    ok = sum(r["Статус"] == "OK" for r in new_rows)
    failed = len(new_rows) - ok
    if new_rows:
        st.success(f"Готово: успешно {ok}, с ошибками {failed}")


def render_results() -> None:
    """Таблица результатов и кнопки экспорта/очистки."""
    if not st.session_state.results:
        st.info("Загрузите PDF и нажмите «Анализировать» — результаты появятся здесь.")
        return

    df = pd.DataFrame(st.session_state.results, columns=RESULT_COLUMNS)

    only_errors = st.checkbox("Показать только ошибки")
    view = df[df["Статус"] == "Ошибка"] if only_errors else df

    st.dataframe(
        view,
        use_container_width=True,
        hide_index=True,
        column_config={
            "Страниц": st.column_config.NumberColumn(format="%d"),
            "Символов": st.column_config.NumberColumn(format="%d"),
            "Краткое содержание": st.column_config.TextColumn(width="large"),
            "Ошибка": st.column_config.TextColumn(width="medium"),
        },
    )

    col_export, col_clear = st.columns([1, 1])
    with col_export:
        try:
            st.download_button(
                "⬇️ Экспорт в Excel",
                data=build_excel(df),
                file_name=f"pdf_analysis_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        except Exception as exc:
            st.error(f"Не удалось сформировать Excel: {exc}")
    with col_clear:
        if st.button("🗑️ Очистить результаты", use_container_width=True):
            st.session_state.results = []
            st.session_state.processed = set()
            st.rerun()


def main() -> None:
    st.set_page_config(page_title="PDF Analyzer", page_icon="📄", layout="wide")
    init_state()

    st.title("📄 Анализ PDF с помощью ИИ")
    st.caption("Загрузите PDF → извлечение текста → анализ моделью → таблица → Excel")

    settings = render_sidebar()

    uploaded = st.file_uploader(
        "PDF-файлы", type=["pdf"], accept_multiple_files=True,
        help="Можно выбрать несколько файлов одновременно",
    )

    if st.button("🚀 Анализировать", type="primary", disabled=not uploaded):
        errors = validate_settings(settings)
        if errors:
            for e in errors:
                st.error(e)
        else:
            process_files(uploaded, settings)

    st.divider()
    render_results()


if __name__ == "__main__":
    main()